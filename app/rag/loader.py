from __future__ import annotations

from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
import re
import shutil
import subprocess
import json

from app.core.config import get_settings
from app.services.ocr_service import OcrService

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass(slots=True)
class LoadedSection:
    text: str
    metadata: dict[str, str | int | float | bool]


@dataclass(slots=True)
class LoadedPage:
    page_number: int
    sections: list[LoadedSection] = field(default_factory=list)
    page_summary: str = ""
    page_type: str = "text"


@dataclass(slots=True)
class LoadedDocument:
    document_id: str
    source_name: str
    source_type: str
    source_path: Path
    sections: list[LoadedSection]
    pages: list[LoadedPage] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(section.text for section in self.sections if section.text.strip())


class DocumentLoader:
    """Load raw files into normalized, structured text sections."""

    heading_pattern = re.compile(r"^(\d+(\.\d+){0,4}|[A-Z][A-Z0-9 ._-]{1,80}|第[一二三四五六七八九十0-9]+[章节部分篇])")
    table_line_pattern = re.compile(r"\s{2,}|\t|\|")
    figure_pattern = re.compile(r"^(图|figure)\s*[\dA-Za-z一二三四五六七八九十]+", re.IGNORECASE)
    footnote_pattern = re.compile(r"^(\[\d+\]|\(\d+\)|注[:：])")
    page_marker_pattern = re.compile(r"^(page|页)\s*[\divxlc]+$", re.IGNORECASE)

    def __init__(self) -> None:
        self.settings = get_settings()
        self.ocr_service = OcrService()

    def load(self, file_path: Path, document_id: str, source_name: str | None = None) -> LoadedDocument:
        suffix = file_path.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            raise ValueError(f"Unsupported file type '{suffix}'. Supported types: {supported}.")

        if suffix == ".pdf":
            pages = self._load_pdf(file_path, document_id=document_id)
        elif suffix == ".docx":
            pages = self._load_docx(file_path)
        else:
            pages = self._load_text(file_path)

        pages = self._filter_repeated_page_noise(pages)

        sections: list[LoadedSection] = []
        for page in pages:
            page.sections = self._enrich_page_sections(page, source_name or file_path.name)
            sections.extend(section for section in page.sections if section.text.strip())

        if not sections:
            raise ValueError(f"No readable text found in '{file_path.name}'.")

        return LoadedDocument(
            document_id=document_id,
            source_name=source_name or file_path.name,
            source_type=suffix.lstrip("."),
            source_path=file_path,
            sections=sections,
            pages=pages,
        )

    def _load_pdf(self, file_path: Path, document_id: str) -> list[LoadedPage]:
        if self._should_use_mineru():
            mineru_pages = self._load_pdf_with_mineru(file_path=file_path, document_id=document_id)
            if mineru_pages:
                return mineru_pages

        from pypdf import PdfReader

        reader = PdfReader(str(file_path))
        pages: list[LoadedPage] = []

        for index, page in enumerate(reader.pages, start=1):
            raw_text = page.extract_text() or ""
            normalized = self._normalize_text(raw_text)
            sections = self._build_page_sections(raw_text, page_number=index)
            page_type = self._detect_pdf_page_type(raw_text=raw_text, sections=sections)

            if page_type == "ocr_candidate":
                ocr_text = self._extract_pdf_page_ocr(file_path=file_path, page_number=index)
                if ocr_text:
                    sections = self._build_page_sections(ocr_text, page_number=index)
                    page_type = "ocr_text"

            image_path = self._export_pdf_page_preview(file_path=file_path, page_number=index)
            if image_path:
                for section in sections:
                    section.metadata["page_image_path"] = str(image_path)

            page_summary = self._build_page_summary(sections)

            pages.append(
                LoadedPage(
                    page_number=index,
                    sections=sections,
                    page_summary=page_summary,
                    page_type=page_type,
                )
            )

        return pages

    def _should_use_mineru(self) -> bool:
        if not self.settings.mineru_enabled:
            return False
        if self.settings.advanced_pdf_backend not in {"auto", "mineru"}:
            return False
        return shutil.which(self.settings.mineru_command) is not None

    def _load_pdf_with_mineru(self, file_path: Path, document_id: str) -> list[LoadedPage]:
        output_root = Path(self.settings.processed_data_dir) / "_mineru" / document_id
        output_root.mkdir(parents=True, exist_ok=True)

        command = self._build_mineru_command(file_path=file_path, output_root=output_root)

        try:
            subprocess.run(
                command,
                check=True,
                capture_output=True,
                text=True,
                timeout=self.settings.mineru_timeout_seconds,
            )
        except Exception:
            return []

        content_path = self._find_mineru_content_list(output_root)
        if not content_path:
            return []

        try:
            content_list = json.loads(content_path.read_text(encoding="utf-8"))
        except Exception:
            return []

        return self._build_pages_from_mineru_content(
            content_list=content_list,
            output_root=content_path.parent,
        )

    def _build_mineru_command(self, file_path: Path, output_root: Path) -> list[str]:
        backend = self.settings.mineru_backend.strip() or "pipeline"
        method = self.settings.mineru_method.strip() or "auto"

        command = [
            self.settings.mineru_command,
            "-p",
            str(file_path),
            "-o",
            str(output_root),
            "-b",
            backend,
            "-m",
            method,
            "-t",
            "true",
            "--client-side-output-generation",
            "true",
        ]

        if self.settings.mineru_api_url.strip():
            command.extend(["--api-url", self.settings.mineru_api_url.strip()])

        if backend == "pipeline" and self.settings.mineru_lang.strip():
            command.extend(["-l", self.settings.mineru_lang.strip()])

        if backend.startswith("hybrid") and self.settings.mineru_effort.strip():
            command.extend(["--effort", self.settings.mineru_effort.strip()])
            if self.settings.figure_vision_enabled:
                command.extend(["--image-analysis", "true"])

        return command

    def _find_mineru_content_list(self, output_root: Path) -> Path | None:
        candidates = list(output_root.rglob("*_content_list.json"))
        if candidates:
            return candidates[0]
        candidates = list(output_root.rglob("*_content_list_v2.json"))
        if candidates:
            return candidates[0]
        return None

    def _build_pages_from_mineru_content(self, content_list: list[dict], output_root: Path) -> list[LoadedPage]:
        grouped: dict[int, list[LoadedSection]] = {}

        for item in content_list:
            if not isinstance(item, dict):
                continue
            page_idx = int(item.get("page_idx", 0))
            page_number = page_idx + 1
            page_sections = grouped.setdefault(page_number, [])
            section = self._mineru_item_to_section(item, output_root=output_root)
            if section is not None:
                page_sections.append(section)

        pages: list[LoadedPage] = []
        for page_number in sorted(grouped):
            sections = grouped[page_number]
            page_type = self._infer_mineru_page_type(sections)
            page_summary = self._build_page_summary(sections)
            pages.append(
                LoadedPage(
                    page_number=page_number,
                    sections=sections,
                    page_summary=page_summary,
                    page_type=page_type,
                )
            )
        return pages

    def _mineru_item_to_section(self, item: dict, output_root: Path) -> LoadedSection | None:
        item_type = str(item.get("type", "text"))
        bbox = item.get("bbox", [])
        page_idx = int(item.get("page_idx", 0))

        if item_type in {"header", "footer", "page_number", "aside_text", "page_footnote"}:
            return None

        metadata: dict[str, str | int | float | bool] = {
            "page": page_idx + 1,
            "section_index": 0,
            "element_type": "paragraph",
            "section_title": "",
            "parent_section": "",
            "heading_level": 0,
            "table_title": "",
            "figure_caption": "",
            "page_type": "mineru",
            "bbox": str(bbox),
        }

        if item_type == "text":
            text = self._normalize_text(str(item.get("text", "")))
            level = int(item.get("text_level", 0) or 0)
            if level > 0:
                metadata["element_type"] = "heading"
                metadata["heading_level"] = level
                metadata["section_title"] = text
            return LoadedSection(text=text, metadata=metadata) if text else None

        if item_type == "list":
            items = item.get("list_items", [])
            text = self._normalize_text("\n".join(str(entry) for entry in items))
            metadata["element_type"] = "paragraph"
            return LoadedSection(text=text, metadata=metadata) if text else None

        if item_type == "table":
            caption = self._normalize_text(" ".join(item.get("table_caption", [])))
            footnote = self._normalize_text(" ".join(item.get("table_footnote", [])))
            body = self._normalize_text(str(item.get("table_body", "")))
            text = "\n".join(part for part in [caption, body, footnote] if part)
            metadata["element_type"] = "table"
            metadata["table_title"] = caption
            img_path = item.get("img_path")
            if isinstance(img_path, str) and img_path:
                metadata["page_image_path"] = str((output_root / img_path).resolve())
            return LoadedSection(text=text, metadata=metadata) if text else None

        if item_type in {"image", "chart"}:
            caption_key = "image_caption" if item_type == "image" else "chart_caption"
            footnote_key = "image_footnote" if item_type == "image" else "chart_footnote"
            content_key = "content"
            caption = self._normalize_text(" ".join(item.get(caption_key, [])))
            footnote = self._normalize_text(" ".join(item.get(footnote_key, [])))
            content = self._normalize_text(str(item.get(content_key, "")))
            text = "\n".join(part for part in [caption, content, footnote] if part)
            metadata["element_type"] = "figure_caption"
            metadata["figure_caption"] = caption or content
            img_path = item.get("img_path")
            if isinstance(img_path, str) and img_path:
                metadata["page_image_path"] = str((output_root / img_path).resolve())
            return LoadedSection(text=text, metadata=metadata) if text else None

        if item_type == "equation":
            text = self._normalize_text(str(item.get("text", "")))
            metadata["element_type"] = "paragraph"
            return LoadedSection(text=text, metadata=metadata) if text else None

        return None

    def _infer_mineru_page_type(self, sections: list[LoadedSection]) -> str:
        element_types = {str(section.metadata.get("element_type", "paragraph")) for section in sections}
        if "table" in element_types:
            return "table_heavy"
        if "figure_caption" in element_types:
            return "figure_mixed"
        return "text"

    def _load_docx(self, file_path: Path) -> list[LoadedPage]:
        from docx import Document

        document = Document(str(file_path))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        raw_text = "\n".join(paragraphs)
        sections = self._build_page_sections(raw_text, page_number=1)
        return [LoadedPage(page_number=1, sections=sections, page_summary=self._build_page_summary(sections), page_type="text")]

    def _load_text(self, file_path: Path) -> list[LoadedPage]:
        raw_text = file_path.read_text(encoding="utf-8", errors="ignore")
        sections = self._build_page_sections(raw_text, page_number=1)
        return [LoadedPage(page_number=1, sections=sections, page_summary=self._build_page_summary(sections), page_type="text")]

    def _build_page_sections(self, text: str, page_number: int) -> list[LoadedSection]:
        if not text:
            return []

        normalized_text = self._normalize_text_preserve_layout(text)
        blocks = [block.strip() for block in re.split(r"\n{2,}", normalized_text) if block.strip()]
        if len(blocks) <= 1 and normalized_text:
            line_blocks = [line.strip() for line in normalized_text.splitlines() if line.strip()]
            if line_blocks:
                blocks = line_blocks
        sections: list[LoadedSection] = []
        current_headings: list[str] = []
        seen_table = 0
        seen_figure = 0

        for block_index, block in enumerate(blocks):
            lines = [line.strip() for line in block.splitlines() if line.strip()]
            if not lines:
                continue

            first_line = lines[0]
            element_type = "paragraph"
            heading_level: int | None = None
            title = ""

            if self._is_table_block(lines):
                element_type = "table"
                seen_table += 1
                title = first_line if self._is_table_title(first_line) else f"Table {seen_table}"
            elif self._is_heading(first_line, lines):
                element_type = "heading"
                heading_level = self._infer_heading_level(first_line)
                title = first_line
                current_headings = current_headings[: max(heading_level - 1, 0)]
                current_headings.append(first_line)
            elif self.figure_pattern.match(first_line):
                element_type = "figure_caption"
                seen_figure += 1
                title = first_line or f"Figure {seen_figure}"
            elif self.footnote_pattern.match(first_line):
                element_type = "footnote"
            elif self._looks_scanned_block(block):
                element_type = "ocr_candidate"

            metadata: dict[str, str | int | float | bool] = {
                "page": page_number,
                "section_index": block_index,
                "element_type": element_type,
                "section_title": current_headings[-1] if current_headings else "",
                "parent_section": " > ".join(current_headings[:-1]) if len(current_headings) > 1 else "",
                "heading_level": heading_level or 0,
                "table_title": title if element_type == "table" else "",
                "figure_caption": title if element_type == "figure_caption" else "",
                "page_type": "text",
            }
            sections.append(LoadedSection(text=block, metadata=metadata))

        return sections

    def _enrich_page_sections(self, page: LoadedPage, source_name: str) -> list[LoadedSection]:
        enriched: list[LoadedSection] = []
        for section in page.sections:
            metadata = dict(section.metadata)
            metadata["source_name"] = source_name
            metadata["page_summary"] = page.page_summary
            metadata["page_type"] = page.page_type
            metadata["page_anchor"] = f"{source_name}#page={page.page_number}"
            enriched.append(LoadedSection(text=section.text, metadata=metadata))

        if page.page_summary:
            enriched.append(
                LoadedSection(
                    text=page.page_summary,
                    metadata={
                        "page": page.page_number,
                        "section_index": len(enriched),
                        "element_type": "page_summary",
                        "section_title": "",
                        "parent_section": "",
                        "heading_level": 0,
                        "table_title": "",
                        "figure_caption": "",
                        "page_type": page.page_type,
                        "page_summary": page.page_summary,
                        "page_anchor": f"{source_name}#page={page.page_number}",
                        "source_name": source_name,
                    },
                )
            )

        return enriched

    def _filter_repeated_page_noise(self, pages: list[LoadedPage]) -> list[LoadedPage]:
        repeated_lines: Counter[str] = Counter()

        for page in pages:
            for section in page.sections:
                for line in section.text.splitlines():
                    cleaned = line.strip()
                    if cleaned and len(cleaned) <= 80:
                        repeated_lines[cleaned] += 1

        noise_candidates = {
            line
            for line, count in repeated_lines.items()
            if count >= max(2, len(pages) // 2)
            and (
                self.page_marker_pattern.match(line)
                or any(token in line.lower() for token in ("confidential", "copyright", "版权所有", "保密", "第"))
            )
        }

        filtered_pages: list[LoadedPage] = []
        for page in pages:
            filtered_sections: list[LoadedSection] = []
            for section in page.sections:
                lines = [line for line in section.text.splitlines() if line.strip() not in noise_candidates]
                text = self._normalize_text("\n".join(lines))
                if not text:
                    continue
                filtered_sections.append(LoadedSection(text=text, metadata=dict(section.metadata)))

            page.sections = filtered_sections
            page.page_summary = self._build_page_summary(filtered_sections)
            filtered_pages.append(page)

        return filtered_pages

    def _build_page_summary(self, sections: list[LoadedSection]) -> str:
        if not sections:
            return ""

        summary_parts: list[str] = []
        seen_types: set[str] = set()

        for section in sections:
            element_type = str(section.metadata.get("element_type", "paragraph"))
            if element_type == "page_summary":
                continue
            if element_type not in seen_types or element_type in {"heading", "table", "figure_caption"}:
                snippet = section.text.replace("\n", " ").strip()
                summary_parts.append(f"{element_type}: {snippet[:100]}")
                seen_types.add(element_type)
            if len(summary_parts) >= 4:
                break

        return " | ".join(summary_parts)

    def _detect_pdf_page_type(self, raw_text: str, sections: list[LoadedSection]) -> str:
        normalized = self._normalize_text(raw_text)
        if not normalized or len(normalized) < 40:
            return "ocr_candidate"

        table_sections = [section for section in sections if section.metadata.get("element_type") == "table"]
        figure_sections = [section for section in sections if section.metadata.get("element_type") == "figure_caption"]

        if table_sections:
            return "table_heavy"
        if figure_sections:
            return "figure_mixed"
        return "text"

    def _extract_pdf_page_ocr(self, file_path: Path, page_number: int) -> str:
        if not self.ocr_service.available:
            return ""

        image_path = self._export_pdf_page_preview(file_path=file_path, page_number=page_number)
        if not image_path:
            return ""

        return self._normalize_text(self.ocr_service.extract_text(image_path))

    def _export_pdf_page_preview(self, file_path: Path, page_number: int) -> Path | None:
        try:
            import fitz
        except Exception:
            return None

        output_dir = file_path.parent.parent / "page_previews"
        output_dir.mkdir(parents=True, exist_ok=True)
        image_path = output_dir / f"{file_path.stem}-page-{page_number:03d}.png"

        if image_path.exists():
            return image_path

        try:
            document = fitz.open(str(file_path))
            page = document.load_page(page_number - 1)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            pixmap.save(str(image_path))
            document.close()
            return image_path
        except Exception:
            return None

    def _is_heading(self, first_line: str, lines: list[str]) -> bool:
        if len(lines) > 2:
            return False
        if self.figure_pattern.match(first_line):
            return False
        if len(first_line) > 100:
            return False
        if self._is_table_title(first_line):
            return True
        return bool(self.heading_pattern.match(first_line))

    def _infer_heading_level(self, text: str) -> int:
        match = re.match(r"^(\d+(\.\d+){0,4})", text)
        if match:
            return text.count(".") + 1
        return 1

    def _is_table_block(self, lines: list[str]) -> bool:
        if not lines:
            return False
        table_like_lines = sum(
            1
            for line in lines[:6]
            if self.table_line_pattern.search(line) or re.search(r"\b\d+(\.\d+)?\b", line)
        )
        return table_like_lines >= max(2, min(len(lines), 3))

    def _is_table_title(self, text: str) -> bool:
        return bool(re.match(r"^(表|table)\s*[\dA-Za-z一二三四五六七八九十]+", text, re.IGNORECASE))

    def _looks_scanned_block(self, text: str) -> bool:
        compact = text.replace(" ", "")
        if len(compact) < 20:
            return True
        non_word_ratio = len(re.findall(r"[^\w\u4e00-\u9fff]", compact)) / max(len(compact), 1)
        return non_word_ratio > 0.35

    def _normalize_text(self, text: str) -> str:
        compact = text.replace("\u0000", " ")
        compact = re.sub(r"\r\n?", "\n", compact)
        compact = re.sub(r"[ \t]+", " ", compact)
        compact = re.sub(r"\n{3,}", "\n\n", compact)
        return compact.strip()

    def _normalize_text_preserve_layout(self, text: str) -> str:
        compact = text.replace("\u0000", " ")
        compact = re.sub(r"\r\n?", "\n", compact)
        compact = re.sub(r"[ ]{4,}", "    ", compact)
        compact = re.sub(r"\t+", "\t", compact)
        compact = re.sub(r"\n{3,}", "\n\n", compact)
        return compact.strip()

    def _load_text_fixture(self, page_number: int, blocks: list[str]) -> LoadedPage:
        text = self._normalize_text("\n\n".join(blocks))
        sections = self._build_page_sections(text, page_number=page_number)
        return LoadedPage(
            page_number=page_number,
            sections=sections,
            page_summary=self._build_page_summary(sections),
            page_type="text",
        )
